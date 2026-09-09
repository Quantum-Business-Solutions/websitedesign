#!/usr/bin/env python3
"""Search and AI-answer package for a preview: the full report, the every-page audit on the hub,
the audit spreadsheet and the redirect map.

Inputs, all next to the content file:
  brands/<slug>.seo.json    the measured facts and the written findings (authored per client)
  brands/<slug>.audit.json  every page of the live site, scored by seo_audit.py

preview.py calls write() after the direction folders exist; the build is scored with the same
sixteen checks so the hub can show today against the build, page for page. Returns the hub
fragments (search section, audit section, CSS, nav link) or None when no seo.json exists.
"""
import csv
import html
import io
import json
import os
import re
import sys

import seo_audit as A

HERE = os.path.dirname(os.path.abspath(__file__))

E = lambda s: html.escape(str(s if s is not None else ""), quote=True)  # noqa: E731

SEV = {"Critical": "#b42318", "High": "#c4320a", "Medium": "#b54708", "Low": "#5b616b"}
GRADE = {"A": "#1b7f4b", "B": "#3b7d3a", "C": "#b54708", "D": "#c4320a", "F": "#b42318"}
DONE = {
    "title_len": "Title cut to 60 characters or fewer", "city": "City added to the title", "meta_len": "Meta description rewritten to length",
    "h1": "One H1 on the page", "h2": "Copy broken into H2 sections", "question": "A question heading added", "words": "Copy deepened past 300 words",
    "alt": "Alt text on every image", "links": "Ten or more internal links", "faq": "FAQ marked up as FAQPage", "service": "Service or Article schema added",
    "local": "LocalBusiness schema added", "review": "Review schema waits for the Google review program", "canonical": "Self-referencing canonical",
    "og": "Branded share image", "form": "Form or tap-to-call number added",
}


REQUIRED = ["measured", "hero", "findings", "compare", "moves", "entity", "today", "competitors", "opportunities", "clusters", "questions",
            "money_pages", "competitor_pages", "aeo", "crawlers", "backlinks", "local", "plan", "measure", "hub", "sitewide"]


def validate(seo, path):
    missing = [k for k in REQUIRED if k not in seo]
    for k in ("today", "after", "keep", "heading", "intro"):
        if "hub" in seo and k not in seo["hub"]:
            missing.append(f"hub.{k}")
    if len(seo.get("findings", [])) < 5:
        missing.append("findings (fewer than five)")
    if len(seo.get("moves", [])) != 5:
        missing.append("moves (needs exactly five)")
    if missing:
        raise SystemExit(f"{path}: the search package is incomplete: {', '.join(missing)}")


def _num(v):
    try:
        return float(str(v).replace(",", "").split(" ")[0])
    except ValueError:
        return None


def _gauge(today, build):
    """Two arcs on one track: today in the muted red, the build in green, the numbers in the middle."""
    import math
    r, c = 54, 70
    circ = 2 * math.pi * r
    def arc(v, color, w):
        return f'<circle cx="{c}" cy="{c}" r="{r}" fill="none" stroke="{color}" stroke-width="{w}" stroke-linecap="round" stroke-dasharray="{circ * v / 100:.1f} {circ:.1f}" transform="rotate(-90 {c} {c})"/>'
    if build is None:  # analysis only: one arc, today's score
        return (f'<svg class="gauge" viewBox="0 0 140 140" role="img" aria-label="Readiness score today {today} of 100">'
                f'<circle cx="{c}" cy="{c}" r="{r}" fill="none" stroke="rgba(255,255,255,.14)" stroke-width="12"/>{arc(today, "#ffb4a8", 12)}'
                f'<text x="{c}" y="{c - 4}" text-anchor="middle" font-size="30" font-weight="800" fill="#fff">{today}</text>'
                f'<text x="{c}" y="{c + 18}" text-anchor="middle" font-size="12" fill="rgba(255,255,255,.8)">of 100 today</text></svg>')
    return (f'<svg class="gauge" viewBox="0 0 140 140" role="img" aria-label="Readiness score today {today} of 100, the build {build} of 100">'
            f'<circle cx="{c}" cy="{c}" r="{r}" fill="none" stroke="rgba(255,255,255,.14)" stroke-width="12"/>{arc(build, "#8fe3b0", 12)}{arc(today, "#ffb4a8", 6)}'
            f'<text x="{c}" y="{c - 4}" text-anchor="middle" font-size="30" font-weight="800" fill="#fff">{build}</text>'
            f'<text x="{c}" y="{c + 18}" text-anchor="middle" font-size="12" fill="rgba(255,255,255,.8)">from {today} today</text></svg>')


def _bars(rows, idx, label, highlight, fmt=lambda v: f"{int(v):,}"):
    """Horizontal bars for one numeric column of a table; the client's row in the accent."""
    data = [(r[0], _num(r[idx])) for r in rows]
    data = [(k, v) for k, v in data if v is not None]
    if not data:
        return ""
    mx = max(v for _, v in data) or 1
    items = "".join(f'<div class="bar-row{" me" if highlight in k else ""}"><span class="bar-k">{E(k.split(" (")[0])}</span><span class="bar-t"><i style="width:{max(2, v / mx * 100):.1f}%"></i></span><span class="bar-v">{E(fmt(v))}</span></div>' for k, v in data)
    return f'<figure class="bars"><figcaption>{E(label)}</figcaption>{items}</figure>'



def _donut(pct, label, sub, color):
    import math
    r, c = 40, 50
    circ = 2 * math.pi * r
    return (f'<figure class="donut" aria-label="{E(label)}: {pct} percent"><svg viewBox="0 0 100 100"><circle cx="{c}" cy="{c}" r="{r}" fill="none" stroke="var(--alt)" stroke-width="14"/>'
            f'<circle cx="{c}" cy="{c}" r="{r}" fill="none" stroke="{color}" stroke-width="14" stroke-dasharray="{circ * pct / 100:.1f} {circ:.1f}" transform="rotate(-90 {c} {c})"/>'
            f'<text x="{c}" y="{c + 8}" text-anchor="middle" font-size="24" font-weight="800" fill="var(--fg)">{pct}%</text></svg><figcaption><b>{E(label)}</b><span>{E(sub)}</span></figcaption></figure>')


def _grades(pages_today, pages_build):
    """Grouped bars: how many pages sit at each grade today and in the build."""
    order = "ABCDF"
    ct = {g: sum(1 for p in pages_today if p["grade"] == g) for g in order}
    if pages_build is None:
        mx = max(ct.values()) or 1
        cols = "".join(f'<div class="gcol"><div class="gbars"><i class="t" style="height:{ct[g] / mx * 100:.0f}%;background:{GRADE[g]}" title="{ct[g]} pages"></i></div><span class="gl"><b style="color:{GRADE[g]}">{g}</b><small>{ct[g]} pages</small></span></div>' for g in order)
        return f'<figure class="grades" aria-label="Pages by grade today"><figcaption>Pages by grade today, {len(pages_today)} pages</figcaption><div class="grow">{cols}</div></figure>'
    cb = {g: sum(1 for p in pages_build if p["grade"] == g) for g in order}
    mx = max(list(ct.values()) + list(cb.values())) or 1
    cols = ""
    for g in order:
        cols += (f'<div class="gcol"><div class="gbars"><i class="t" style="height:{ct[g] / mx * 100:.0f}%" title="today {ct[g]}"></i><i class="b" style="height:{cb[g] / mx * 100:.0f}%" title="build {cb[g]}"></i></div>'
                 f'<span class="gl"><b style="color:{GRADE[g]}">{g}</b><small>{ct[g]} today<br>{cb[g]} build</small></span></div>')
    return f'<figure class="grades" aria-label="Pages by grade, today against the build"><figcaption>Pages by grade, today against the build <span class="key"><i class="t"></i> today <i class="b"></i> the build</span></figcaption><div class="grow">{cols}</div></figure>'


def _scatter(rows, highlight):
    """Authority Score against organic visits, one dot per competitor, the client in the accent.
    Linear scales from zero with ticks; labels swap side on the right half and are nudged apart when they collide."""
    pts = []
    for r in rows:
        v, a = _num(r[1]), _num(r[3]) if len(r) > 3 else None
        if v is None or a is None:
            continue
        pts.append((r[0].split(" (")[0], v, a, highlight in r[0]))
    if len(pts) < 2:
        return ""
    import math
    W, H, pad, top = 560, 280, 56, 20

    def nice(m):  # axis top: four equal ticks landing on round numbers
        m = max(m * 1.1, 1)
        unit = 10 ** math.floor(math.log10(m / 4))
        step = next(u for u in (unit, unit * 2, unit * 2.5, unit * 5, unit * 10) if u * 4 >= m)
        return step * 4
    mv = nice(max(p[1] for p in pts))
    ma = nice(max(p[2] for p in pts))
    grid = ""
    for i in range(1, 4):
        yy = H - pad - (i / 4) * (H - pad - top)
        grid += f'<line x1="{pad}" y1="{yy:.1f}" x2="{W - 12}" y2="{yy:.1f}" stroke="var(--line)" stroke-dasharray="3 4"/><text x="{pad - 6}" y="{yy + 4:.1f}" text-anchor="end" font-size="11" fill="var(--muted)">{int(mv * i / 4):,}</text>'
    for i in range(1, 5):
        xx = pad + (i / 4) * (W - pad - 12)
        grid += f'<text x="{xx:.1f}" y="{H - pad + 16}" text-anchor="middle" font-size="11" fill="var(--muted)">{int(ma * i / 4)}</text>'
    placed = []  # (x, y, label_y, side, name, me)

    def span(x, side, name):  # horizontal extent of a label, about 6.6 px a character at 12 px
        w = 6.6 * len(name) + 4
        lx = x + side * 12
        return (lx, lx + w) if side > 0 else (lx - w, lx)

    def clear(ly, sp):
        return all(abs(ly - q[2]) >= 15 or sp[1] < span(q[0], q[3], q[4])[0] or sp[0] > span(q[0], q[3], q[4])[1] for q in placed)
    for name, v, a, me in sorted(pts, key=lambda p: p[1], reverse=True):
        x = pad + (a / ma) * (W - pad - 12)
        y = H - pad - (v / mv) * (H - pad - top)
        side = -1 if x > W * 0.62 else 1
        sp = span(x, side, name)
        ly = y
        for d in (0, 15, -15, 30, -30, 45, -45, 60, -60):  # nudge below, then above, staying inside the plot
            cand = y + d
            if top + 8 <= cand <= H - pad - 4 and clear(cand, sp):
                ly = cand
                break
        placed.append((x, y, ly, side, name, me))
    dots = ""
    for x, y, ly, side, name, me in placed:
        lx = x + side * 12
        if abs(ly - y) > 6:
            dots += f'<line x1="{x:.1f}" y1="{y:.1f}" x2="{lx - side * 3:.1f}" y2="{ly - 3:.1f}" stroke="var(--muted)" stroke-width="0.8"/>'
        dots += f'<circle cx="{x:.1f}" cy="{y:.1f}" r="{9 if me else 6}" fill="{"var(--accent)" if me else "#9aa4b2"}"/><text x="{lx:.1f}" y="{ly + 4:.1f}" font-size="12" text-anchor="{"end" if side < 0 else "start"}" fill="var(--fg)" font-weight="{700 if me else 400}">{E(name)}</text>'
    return (f'<figure class="scatter" aria-label="Authority Score against organic visits a month for the client and its competitors"><figcaption>Authority against visits: who converts links into traffic</figcaption>'
            f'<svg viewBox="0 0 {W} {H}">{grid}<line x1="{pad}" y1="{H - pad}" x2="{W - 12}" y2="{H - pad}" stroke="var(--border)"/><line x1="{pad}" y1="{top}" x2="{pad}" y2="{H - pad}" stroke="var(--border)"/>'
            f'<text x="{(W + pad) / 2}" y="{H - 8}" text-anchor="middle" font-size="12" fill="var(--muted)">Authority Score</text><text x="14" y="{(H - pad + top) / 2}" text-anchor="middle" font-size="12" fill="var(--muted)" transform="rotate(-90 14 {(H - pad + top) / 2})">Organic visits a month</text>{dots}</svg></figure>')


def _opps(rows, short):
    """Volume bars with the client's position as a tag: the gap made visible."""
    data = [(r[0], _num(r[1]), str(r[3])) for r in rows if _num(r[1])]
    data = sorted(data, key=lambda x: -x[1])[:12]
    if not data:
        return ""
    mx = data[0][1]
    items = ""
    for kw, v, pos in data:
        ranking = pos and not pos.lower().startswith("not")
        tag = f'<b class="ok">#{E(pos)}</b>' if ranking and pos.replace(",", "").isdigit() else (f'<b class="ok">{E(pos)}</b>' if ranking else '<b class="bad">not ranking</b>')
        items += f'<div class="bar-row"><span class="bar-k">{E(kw)}</span><span class="bar-t"><i style="width:{max(2, v / mx * 100):.1f}%"></i></span><span class="bar-v">{int(v):,}</span><span class="bar-p">{tag}</span></div>'
    return f'<figure class="bars opps"><figcaption>Searches a month, and where {E(short)} stands</figcaption>{items}</figure>'


def _aeo_matrix(rows, short):
    """Ten cells per query: the client's position lit, everything else grey. Presence at a glance."""
    out = ""
    for r in rows:
        pos = r.get("position")
        cells = "".join(f'<i class="{"me" if pos == i else ""}" title="{E(r["results"][i - 1]) if i - 1 < len(r["results"]) else ""}"></i>' for i in range(1, 11))
        out += f'<div class="mrow"><span class="mq">{E(r["query"])}</span><span class="mcells">{cells}</span><span class="mv">{("#" + str(pos)) if pos else "absent"}</span></div>'
    return f'<figure class="matrix" aria-label="Answer-engine scoreboard: {E(short)} position in the top ten for each buyer query"><figcaption>The top ten for each query; the lit cell is {E(short)}</figcaption><div class="mhead"><span></span><span class="mcells">{"".join(f"<em>{i}</em>" for i in range(1, 11))}</span><span></span></div>{out}</figure>'


def _bytype(summary):
    bt = summary.get("by_type") or {}
    if not bt:
        return ""
    rows = sorted(bt.items(), key=lambda kv: -kv[1]["count"])
    items = "".join(f'<div class="bar-row"><span class="bar-k">{E(k.title())} pages ({v["count"]})</span><span class="bar-t"><i style="width:{v["avg"]}%;background:{GRADE["A" if v["avg"] >= 85 else "B" if v["avg"] >= 70 else "C" if v["avg"] >= 55 else "D" if v["avg"] >= 40 else "F"]}"></i></span><span class="bar-v">{v["avg"]}</span></div>' for k, v in rows)
    return f'<figure class="bars"><figcaption>Average readiness by page type, out of 100</figcaption>{items}</figure>'


def _load(content_path):
    seo_p = content_path.replace(".content.json", ".seo.json")
    aud_p = content_path.replace(".content.json", ".audit.json")
    if not (os.path.exists(seo_p) and os.path.exists(aud_p)):
        return None, None
    return json.load(open(seo_p, encoding="utf-8")), json.load(open(aud_p, encoding="utf-8"))


def _tokens(s):
    return {t for t in re.split(r"[^a-z0-9]+", s.lower()) if t and t not in ("the", "and", "a", "an", "of", "for", "in", "to", "html", "www", "com", "us", "your", "our", "how", "what", "why", "is", "it")}


def map_pages(live, build_files, seo):
    """old path -> (new file, kind). kind is rebuilt, merged or retired."""
    over = seo.get("redirect_overrides", {})
    retire = [re.compile(r) for r in seo.get("retire", [])]
    rules = [(re.compile(rx), tgt) for rx, tgt in seo.get("redirect_rules", [])]
    stems = {f: _tokens(f) for f in build_files}
    out = {}
    for p in live:
        path = p["path"]
        if path in over:
            out[path] = (over[path], "rebuilt" if over[path] not in ("index.html", "blog.html") else "merged")
            continue
        if any(r.search(path) for r in retire):
            out[path] = (seo.get("retire_target", "index.html"), "retired")
            continue
        hit = None
        for rx, tgt in rules:
            m = rx.search(path)
            if m:
                cand = rx.sub(tgt, path) if "\\" in tgt else tgt
                cand = cand.lstrip("/")
                if cand in build_files:
                    hit = (cand, "rebuilt" if "*" not in tgt and "\\" in tgt else "merged")
                else:
                    hit = (tgt.lstrip("/"), "merged") if tgt.lstrip("/") in build_files else None
                if hit:
                    break
        if hit:
            out[path] = hit
            continue
        toks = _tokens(path.rstrip("/").split("/")[-1]) | _tokens(p.get("h1", ""))
        best, score = None, 0.0
        for f, st in stems.items():
            if not st or not toks:
                continue
            j = len(toks & st) / len(toks | st)
            if j > score:
                best, score = f, j
        if best and score >= 0.34:
            out[path] = (best, "rebuilt")
        elif p.get("type") == "post" and seo.get("post_target"):
            out[path] = (seo["post_target"], "migrated")
        else:
            out[path] = (seo.get("merge_target", "index.html"), "merged")
    return out


def _link(url, text=None, cls=""):
    """An anchor to a live or build page, opening in its own tab so the report stays put."""
    k = f' class="{cls}"' if cls else ""
    return f'<a href="{E(url)}" target="_blank" rel="noopener"{k}>{E(text if text is not None else url)}</a>'


def _tbl(head, rows, cls="", scope_row=False):
    th = "".join(f'<th scope="col">{E(h)}</th>' for h in head)
    body = []
    for r in rows:
        cells = []
        for i, c in enumerate(r):
            tag = "th" if (scope_row and i == 0) else "td"
            kls = ""
            if isinstance(c, tuple):
                c, kls = c
                kls = f' class="{kls}"'
            sc = ' scope="row"' if tag == "th" else ""
            cells.append(f"<{tag}{sc}{kls}>{c if isinstance(c, str) and c.startswith('<') else E(c)}</{tag}>")
        body.append("<tr>" + "".join(cells) + "</tr>")
    return f'<div class="tblwrap" tabindex="0"><table class="{cls}"><thead><tr>{th}</tr></thead><tbody>{"".join(body)}</tbody></table></div>'


def _stiles(tiles, cls=""):
    return f'<div class="stiles {cls}">' + "".join(f'<div class="stile"><b>{E(v)}</b><span>{E(l)}</span></div>' for v, l in tiles) + "</div>"


def _finding_cards(findings, cls="finding"):
    out = []
    for i, f in enumerate(findings, 1):
        sev = f.get("severity", "Medium")
        out.append(f'<article class="{cls}"><div class="rank">{i}</div><div><span class="sev" style="--sev:{SEV.get(sev, "#5b616b")}">{E(sev)}</span><h3>{E(f["title"])}</h3>'
                   f'<dl><div><dt>Evidence</dt><dd>{E(f["evidence"])}</dd></div><div><dt>What it costs</dt><dd>{E(f["cost"])}</dd></div><div><dt>What we do</dt><dd>{E(f["fix"])}</dd></div></dl></div></article>')
    return "".join(out)


def _cmp_rows(today, build, n_today, n_build):
    def frac(a, n):
        return f"{a} of {n}"
    rows = [
        ("Average readiness score", f"{today['avg_score']} of 100", f"{build['avg_score']} of 100"),
        ("Pages scoring A", frac(today["grade_a"], n_today), frac(build["grade_a"], n_build)),
        ("Pages scoring D or F", frac(today["grade_d_f"], n_today), frac(build["grade_d_f"], n_build)),
        ("Pages with FAQPage schema", frac(today["faq"], n_today), frac(build["faq"], n_build)),
        ("Pages with Service or Article schema", frac(today["service_or_article"], n_today), frac(build["service_or_article"], n_build)),
        ("Pages with LocalBusiness schema", frac(today["local"], n_today), frac(build["local"], n_build)),
        ("Pages with a question-form heading", frac(today["question"], n_today), frac(build["question"], n_build)),
        ("Pages under 300 words", frac(today["under_300"], n_today), frac(build["under_300"], n_build)),
        ("Titles with the city in them", frac(today["city_title"], n_today), frac(build["city_title"], n_build)),
        ("Pages with an image missing alt text", frac(today["no_alt_pages"], n_today), frac(build["no_alt_pages"], n_build)),
    ]
    return rows


def report_html(content, seo, audit, build, cmp_rows, base, dslugs, roles_avg, build_by_path, mapping):
    b = content["brand"]
    client = content["client"]
    dom = audit["domain"]
    date = seo.get("measured") or audit.get("measured", "")
    import reskin
    accent = b["accent"]
    ink = reskin.darken_until(accent, "#f3f6f9")  # the alternating section tint, the harder of the two grounds ink sits on
    tint = reskin.rgb_to_hex(*(a * 0.14 + 0.86 for a in reskin.hex_to_rgb(accent)))  # the 14 percent tint the tags sit on
    ink2 = reskin.darken_until(accent, tint)
    navy = b.get("chrome_bg") or "#0f1e33"
    h = seo["hero"]
    solo = build is None
    gauge = _gauge(audit["summary"]["avg_score"], None if solo else build["summary"]["avg_score"])
    every = _every_section("every", audit, seo, build_by_path, mapping, cmp_rows, roles_avg, dom, date)
    comp_bars = _bars(seo["competitors"]["rows"], 1, "Organic visits a month", dom) + _scatter(seo["competitors"]["rows"], dom)
    short = client.split(" ")[0]
    tr = seo.get("traffic") or {}
    donuts = ""
    if tr:
        donuts = '<div class="donuts">' + _donut(tr.get("home_share", 0), "Visits landing on the home page", tr.get("home_note", "one page carries the site"), "#ffb4a8") + _donut(tr.get("brand_share", 0), "Visits from the brand name", tr.get("brand_note", "people who already know the company"), "#ffb4a8") + _donut(tr.get("nonbrand_target", 40), "Non-brand share we build toward", "by day 90, measured against this baseline", "#8fe3b0") + "</div>"
    grades = _grades(audit["pages"], None if solo else build["pages"])
    opps = _opps(seo["opportunities"]["rows"], short)
    matrix = _aeo_matrix(seo["aeo"]["rows"], short)
    bytype = _bytype(audit["summary"])
    auth_bars = _bars(seo["backlinks"]["rows"], 1, "Authority Score", dom, fmt=lambda v: str(int(v)))
    fnd = seo["findings"]
    sev_counts = {}
    for f in fnd:
        sev_counts[f["severity"]] = sev_counts.get(f["severity"], 0) + 1
    chips = "".join(f'<span class="chip" style="--sev:{SEV[k]}">{v} {k.lower()}</span>' for k, v in sev_counts.items())

    cmp = seo["compare"]
    cmp_tbl = _tbl(["Signal"] + cmp["columns"], [[r[0]] + [(c, "bad" if str(c).lower().startswith(("no", "none", "absent", "not")) else ("ok" if str(c).lower().startswith("yes") else "")) for c in r[1:]] for r in cmp["rows"]], "tbl cmp", scope_row=True)

    moves = "".join(f'<article class="levi"><div class="n">{i}</div><div><h3>{E(m["title"])}</h3><div class="meta"><span class="when">{E(m["when"])}</span><span class="tag">{E(m["where"])}</span></div><p>{E(m["body"])}</p></div></article>'
                    for i, m in enumerate(seo["moves"], 1))

    s = audit["summary"]
    n = audit["count"]
    site_tiles = _stiles([(s["avg_score"], "average readiness score out of 100"), (f"{s['grade_b_plus']} of {n}", "pages scoring B or better"), (f"{s['grade_d_f']} of {n}", "pages scoring D or F"),
                          (f"{s['faq']} of {n}", "pages with FAQPage schema"), (f"{s['service_or_article']} of {n}", "pages with Service or Article schema"), (s["review"], "pages with Review or rating schema")], "six")
    # by section
    sections = {}
    for p in audit["pages"]:
        parts = p["path"].strip("/").split("/")
        key = "(top level)" if len(parts) <= 1 else "/" + parts[0] + "/"
        d = sections.setdefault(key, {"n": 0, "faq": 0, "svc": 0, "thin": 0, "words": 0})
        d["n"] += 1; d["faq"] += p["faq"]; d["svc"] += p["service"] or p["article"]; d["thin"] += p["words"] < 300; d["words"] += p["words"]
    sec_rows = [[k, v["n"], v["faq"], v["svc"], v["thin"], f"{round(v['words'] / v['n']):,}"] for k, v in sorted(sections.items(), key=lambda kv: -kv[1]["n"])]
    sec_tbl = _tbl(["Section", "Pages", "With FAQPage", "With Service", "Under 300 words", "Average words"], sec_rows)
    money = [p for p in audit["pages"] if p["type"] in ("service", "industry", "city")]
    money.sort(key=lambda p: p["score"])
    worst_rows = [[_link(p["url"], p["path"]), (f"<span class='sc' style='--g:{GRADE[p['grade']]}'>{p['score']} {p['grade']}</span>", ""), ("yes", "ok") if p["faq"] else (("shown, no schema", "warn") if p.get("faq_visible") else ("no", "bad")), ("yes", "ok") if p["service"] else ("no", "bad"),
                   ("yes", "ok") if p["question_headings"] else ("no", "bad"), f"{p['words']:,}", ("yes", "ok") if p["checks"]["city"] else ("no", "bad"), "; ".join(p["recommendations"][:2])] for p in money[:24]]
    worst_tbl = _tbl(["Service, industry and location page", "Score", "FAQPage", "Service schema", "Question heading", "Words", "City in title", "First two fixes"], worst_rows)

    def found_do(block):
        found = "".join(f"<li>{E(x)}</li>" for x in block.get("found", []))
        do = "".join(f'<p><span class="when">{E(w)}</span> {E(t)}</p>' for w, t in block.get("do", []))
        return f'<h3>What we found</h3><ul>{found}</ul><h3>What we would do</h3><div class="dolist">{do}</div>'

    blog = seo.get("blog")
    blog_html = ""
    if blog:
        blog_html = f'<section><div class="wrap"><p class="eyebrow">The blog</p><h2>{E(blog["heading"])}</h2><p class="lead">{E(blog["intro"])}</p>{_stiles(blog["tiles"])}{found_do(blog)}</div></section>'
    ent = seo["entity"]
    ent_html = f'<section><div class="wrap"><p class="eyebrow">Social and the entity</p><h2>{E(ent["heading"])}</h2><p class="lead">{E(ent["intro"])}</p>{_stiles(ent["tiles"])}{found_do(ent)}</div></section>'

    today_li = "".join(f"<li>{E(x)}</li>" for x in seo["today"])
    comp_tbl = _tbl(["Domain", "Organic visits a month", "Keywords", "Authority Score"], seo["competitors"]["rows"])
    opp_tbl = _tbl(["Keyword", "Searches a month", "Difficulty", f"{content.get('client', '').split(' ')[0]} today", "Who ranks"], [[r[0], r[1], r[2], (r[3], "bad" if str(r[3]).startswith("not") else "ok"), r[4]] for r in seo["opportunities"]["rows"]])
    clusters = "".join(f'<h3>{E(c["name"])} <small>seed: {E(c["seed"])}</small></h3>' + _tbl(["Keyword", "Searches a month", "Difficulty", "CPC", "Intent"], c["rows"]) for c in seo["clusters"])
    qcols = "".join(f'<div class="qcol"><h3>{E(q["topic"])}</h3><ul>' + "".join(f'<li>{E(a)} <span class="vol">{E(v)} a month</span></li>' for a, v in q["items"]) + "</ul></div>" for q in seo["questions"])

    by_path = {p["path"]: p for p in audit["pages"]}
    crawl_rows = []
    for path in seo.get("money_pages", []):
        p = by_path.get(path)
        if not p:
            continue
        crawl_rows.append([_link(f"https://{dom}{path}", path), f"{p['title']} ({p['title_len']} chars)", ("yes", "ok") if p["checks"]["city"] else ("no", "bad"), p["meta_len"] or "none", p["h1"] or "none", p["h2_count"],
                           ", ".join(t for t in p["schema"] if t not in ("ListItem", "EntryPoint", "ImageObject", "PropertyValueSpecification", "SearchAction", "ReadAction", "WebPage", "WebSite", "BreadcrumbList")) or "none",
                           ("yes", "ok") if p["faq"] else ("no", "bad"), f"{p['words']:,}"])
    crawl_tbl = _tbl(["Page", "Title", "City in title", "Meta length", "H1", "H2s", "Schema", "FAQ", "Words"], crawl_rows)
    cp_rows = [[r["domain"], _link(r["url"], r["url"].replace("https://", "").replace("http://", "")), r["title"], f"{r['words']:,}", r["schema"], ("yes", "ok") if r["faq"] else ("no", "bad"), r.get("pricing", "no"), ("yes", "ok") if r.get("reviews") else ("no", "bad"), r.get("form", "no")] for r in seo.get("competitor_pages", [])]
    cp_tbl = _tbl(["Competitor", "Page", "Title", "Words", "Schema", "FAQ", "Pricing", "Reviews", "Form above fold"], cp_rows)

    aeo_rows = [[r["query"], ", ".join(r["results"]), ("yes", "ok") if r.get("position") else ("no", "bad"), r.get("position") or ""] for r in seo["aeo"]["rows"]]
    aeo_tbl = _tbl(["Query", "Top results", f"{client.split(' ')[0]} present", "Position"], aeo_rows)
    crawlers = "".join(f"<li>{E(k)}: <b class='{'ok' if v else 'bad'}'>{'yes' if v else 'no'}</b></li>" for k, v in seo["crawlers"])
    bl = seo["backlinks"]
    bl_tbl = _tbl(["Domain", "Authority Score", "Referring domains", "Backlinks"], bl["rows"])
    ref_tbl = _tbl(["Referring domain", "Domain score"], bl["top_ref"])
    local_li = "".join(f"<li>{E(x)}</li>" for x in seo["local"])
    plan = "".join(f'<div class="step"><div class="when">{E(w)}</div><p>{E(t)}</p></div>' for w, t in seo["plan"])
    measure_tbl = _tbl(["What we measure", "How", "What good looks like"], seo["measure"])
    one_list = "".join(f"<li><b>{E(f['title'])}.</b> {E(f['fix'])}</li>" for f in fnd)

    cmp_box = "".join(f'<div class="cmp-row"><span class="lbl">{E(l)}</span><span class="today">{E(t)}</span><span class="arrow" aria-hidden="true"></span><span class="build">{E(bd)}</span></div>' for l, t, bd in cmp_rows)
    opts = "".join(f'<span class="opt"><b>{E(d.title())}</b> {E(v)}</span>' for d, v in roles_avg)

    cmp_section = "" if solo else f'''<div class="cmpbox"><h3>The same sixteen checks, run on the build</h3><p>Every recommendation in this report is already in place on all three directions. The right-hand column is the build, scored by the same script.</p><div class="cmp-rows"><div class="cmp-row head"><span class="lbl"></span><span class="today">{E(dom)} today</span><span class="arrow"></span><span class="build">The build, any direction</span></div>{cmp_box}</div><div class="opts">{opts}</div><p class="fine" style="color:rgba(255,255,255,.75)">What is deliberately not marked up: Review and AggregateRating. Those wait for the Google review program, because marking up self-published testimonials is against Google's guidelines. That is the gap between the build and 100.</p></div>'''
    redir_note = "" if solo else '; the redirect map is <a href="redirects.csv">redirects.csv</a>'
    return f'''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="robots" content="noindex,nofollow">
<title>SEO and AI search analysis of {E(dom)} | prepared for {E(client)}</title>
<style>:root{{--accent:{accent};--ink:{ink};--ink2:{ink2};--fg:#1c1f24;--muted:#5b616b;--line:#dfe6ee;--border:#dfe6ee;--alt:#f3f6f9;--bg:#ffffff;--bg-alt:#f3f6f9;--navy:{navy};--chrome:{navy}}}*{{box-sizing:border-box}}body{{margin:0;font:16px/1.6 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,Helvetica,Arial,sans-serif;color:var(--fg);background:#fff}}
.wrap{{max-width:1040px;margin:0 auto;padding:0 24px}}section{{padding:56px 0;border-top:1px solid var(--line)}}section:nth-of-type(even){{background:var(--alt)}}
.top{{display:flex;justify-content:space-between;align-items:center;gap:16px;border-bottom:1px solid var(--line);padding:14px 24px;flex-wrap:wrap}}.top a{{text-decoration:none;font-weight:600;color:var(--ink)}}
.eyebrow{{font-size:13px;letter-spacing:.14em;text-transform:uppercase;color:var(--ink);font-weight:700;margin:0 0 12px}}h1{{font-size:clamp(34px,4.4vw,52px);line-height:1.05;letter-spacing:-.02em;margin:0 0 18px;text-wrap:balance}}h2{{font-size:30px;line-height:1.15;letter-spacing:-.01em;margin:0 0 14px;text-wrap:balance}}h3{{font-size:19px;margin:30px 0 8px}}h3 small{{font-weight:400;color:var(--muted);font-size:14px;margin-left:8px}}
.lead{{font-size:18px;color:var(--muted);max-width:760px;margin:0 0 22px}}.hero{{background:var(--navy);color:#fff;padding:64px 0;border-top:0}}.hero .eyebrow{{color:#fff;opacity:.85}}.hero .lead{{color:rgba(255,255,255,.85)}}.hero h1{{color:#fff}}
.stiles{{display:grid;grid-template-columns:repeat(4,1fr);gap:14px;margin-top:28px}}.stiles.six{{grid-template-columns:repeat(3,1fr)}}.stile{{background:#fff;border:1px solid var(--line);border-top:3px solid var(--accent);border-radius:12px;padding:18px 20px}}.stile b{{display:block;font-size:34px;line-height:1;letter-spacing:-.02em;color:var(--ink)}}.stile span{{display:block;margin-top:8px;font-size:13.5px;color:var(--muted)}}
.hero .stile{{background:rgba(255,255,255,.08);border:1px solid rgba(255,255,255,.16);border-top:3px solid var(--accent)}}.hero .stile b{{color:#fff}}.hero .stile span{{color:rgba(255,255,255,.8)}}
.chips{{display:flex;gap:8px;flex-wrap:wrap;margin:0 0 22px}}.chip,.sev{{display:inline-block;font-size:12.5px;font-weight:700;letter-spacing:.06em;text-transform:uppercase;color:#fff;background:var(--sev);border-radius:999px;padding:4px 10px}}
.findings{{display:grid;gap:14px}}.finding{{display:grid;grid-template-columns:56px 1fr;gap:16px;background:#fff;border:1px solid var(--line);border-left:5px solid var(--accent);border-radius:12px;padding:20px 22px;box-shadow:0 10px 30px rgba(0,0,0,.05)}}.finding .rank{{font-size:34px;font-weight:800;color:var(--ink);line-height:1}}.finding h3{{margin:8px 0 10px;font-size:20px}}
.finding dl{{display:grid;grid-template-columns:repeat(3,1fr);gap:14px;margin:0}}.finding dt{{font-size:12px;letter-spacing:.1em;text-transform:uppercase;color:var(--muted);font-weight:700}}.finding dd{{margin:4px 0 0;font-size:15px}}
.tblwrap{{overflow-x:auto;margin:14px 0 6px}}table{{width:100%;border-collapse:collapse;font-size:14.5px;background:#fff}}th,td{{text-align:left;padding:9px 11px;border-bottom:1px solid var(--line);vertical-align:top}}th{{font-size:12px;letter-spacing:.1em;text-transform:uppercase;color:var(--muted);background:#f6f8fa}}code{{font-size:13px;background:#f3f4f6;padding:1px 5px;border-radius:4px}}
.ok{{color:#067647;font-weight:700}}.bad{{color:#b42318;font-weight:700}}.na{{color:var(--muted)}}.vol{{color:var(--muted);font-size:13px}}
.qgrid{{display:grid;grid-template-columns:repeat(2,1fr);gap:22px}}.qcol h3{{margin-top:8px}}ul{{padding-left:18px}}li{{margin:6px 0}}
.steps{{display:grid;grid-template-columns:repeat(5,1fr);gap:14px}}.step{{background:#fff;border:1px solid var(--line);border-top:3px solid var(--accent);border-radius:12px;padding:16px 18px}}.step .when{{font-weight:700;color:var(--ink);margin-bottom:6px}}.step p{{margin:0;font-size:14.5px}}
.tbl.cmp th[scope=row]{{text-align:left;font-weight:600;background:#fff;font-size:14.5px;letter-spacing:0;text-transform:none;color:var(--ink)}}.tbl.cmp td:nth-child(2){{background:color-mix(in srgb,var(--accent) 8%,#fff)}}
.sc{{display:inline-block;color:#fff;background:var(--g);border-radius:6px;padding:2px 8px;font-size:13px;font-weight:700}}
.verdict{{margin-top:22px;padding:22px 26px;background:var(--navy);color:#fff;border-radius:12px;font-size:18px;line-height:1.5}}
.lev{{display:grid;gap:14px}}.levi{{display:grid;grid-template-columns:56px 1fr;gap:16px;background:#fff;border:1px solid var(--line);border-radius:12px;padding:20px 22px;box-shadow:0 10px 30px rgba(0,0,0,.05)}}.levi .n{{font-size:34px;font-weight:800;color:var(--ink);line-height:1}}.levi h3{{margin:6px 0 6px;font-size:20px}}.levi .meta{{display:flex;gap:10px;flex-wrap:wrap;margin:0 0 8px;color:var(--muted);font-size:14px;align-items:center}}.levi .tag{{background:color-mix(in srgb,var(--accent) 14%,#fff);color:var(--ink2);border-radius:999px;padding:2px 10px;font-weight:600}}.levi .when{{font-weight:700;color:var(--ink)}}.levi p{{margin:0}}
.dolist p{{margin:8px 0;padding-left:0}}.dolist .when{{display:inline-block;min-width:78px;font-size:12px;font-weight:700;letter-spacing:.06em;text-transform:uppercase;color:var(--ink)}}
.note{{background:#fff;border-left:4px solid var(--accent);padding:12px 18px;color:var(--muted);margin:24px 0 0}}.fine{{font-size:13px;color:var(--muted);margin-top:14px}}
.btn{{display:inline-block;background:var(--accent);color:#000;padding:12px 20px;border-radius:999px;font-weight:600;text-decoration:none}}
.cmpbox{{margin-top:28px;background:var(--navy);color:#fff;border-radius:16px;padding:26px 28px}}.cmpbox h3{{color:#fff;margin:0 0 6px}}.cmpbox p{{color:rgba(255,255,255,.8);margin:0 0 16px}}.cmp-rows{{display:grid;gap:6px}}.cmp-row{{display:grid;grid-template-columns:1.6fr 1fr 40px 1fr;align-items:center;gap:10px;padding:9px 12px;border-radius:10px;background:rgba(255,255,255,.06)}}.cmp-row.head{{background:none;font-size:12px;letter-spacing:.08em;text-transform:uppercase;color:rgba(255,255,255,.7);padding-bottom:0}}.cmp-row .lbl{{font-weight:600}}.cmp-row .today{{color:#ffb4a8;font-weight:700;text-align:right}}.cmp-row .build{{color:#8fe3b0;font-weight:800;font-size:17px}}.cmp-row .arrow{{width:40px;height:2px;background:rgba(255,255,255,.35);justify-self:center;position:relative}}.cmp-row .arrow::after{{content:"";position:absolute;right:-1px;top:-4px;border:5px solid transparent;border-left-color:rgba(255,255,255,.35)}}.cmp-row.head .arrow{{background:none}}.cmp-row.head .arrow::after{{display:none}}
.opts{{display:flex;gap:10px;flex-wrap:wrap;margin-top:14px}}.opt{{background:rgba(255,255,255,.1);border-radius:999px;padding:6px 12px;font-size:13.5px}}.opt b{{margin-right:4px}}
.hero-grid{{display:grid;grid-template-columns:1fr 240px;gap:40px;align-items:center}}.gauge-wrap{{text-align:center}}.gauge{{width:200px;height:200px}}.gauge-wrap p{{font-size:13px;color:rgba(255,255,255,.75);margin:8px 0 0}}
.bars{{margin:18px 0 26px;padding:18px 20px;background:#fff;border:1px solid var(--line);border-radius:12px}}.bars figcaption{{font-size:12px;letter-spacing:.1em;text-transform:uppercase;color:var(--muted);font-weight:700;margin-bottom:10px}}.bar-row{{display:grid;grid-template-columns:220px 1fr 70px;gap:12px;align-items:center;padding:5px 0;font-size:14px}}.bar-t{{height:10px;background:var(--alt);border-radius:999px;overflow:hidden}}.bar-t i{{display:block;height:100%;background:#9aa4b2;border-radius:999px}}.bar-row.me .bar-t i{{background:var(--accent)}}.bar-row.me .bar-k{{font-weight:700}}.bar-v{{text-align:right;font-variant-numeric:tabular-nums;color:var(--muted)}}
.donut-band{{padding:28px 0;background:#fff;border-top:0}}.donuts{{display:grid;grid-template-columns:repeat(3,1fr);gap:20px}}.donut{{margin:0;display:grid;grid-template-columns:110px 1fr;gap:14px;align-items:center;padding:14px 16px;border:1px solid var(--line);border-radius:12px}}.donut svg{{width:110px;height:110px}}.donut figcaption b{{display:block;font-size:15px}}.donut figcaption span{{display:block;font-size:13px;color:var(--muted);margin-top:4px}}
.grades{{margin:22px 0;padding:18px 20px;background:#fff;border:1px solid var(--line);border-radius:12px}}.grades figcaption,.scatter figcaption,.matrix figcaption{{font-size:12px;letter-spacing:.1em;text-transform:uppercase;color:var(--muted);font-weight:700;margin-bottom:12px;display:flex;justify-content:space-between;gap:12px;flex-wrap:wrap}}.grades .key i{{display:inline-block;width:12px;height:12px;border-radius:3px;vertical-align:middle;margin:0 4px 0 10px}}.grades i.t{{background:#ffb4a8}}.grades i.b{{background:#8fe3b0}}.grow{{display:grid;grid-template-columns:repeat(5,1fr);gap:18px;align-items:end}}.gbars{{height:150px;display:flex;align-items:flex-end;gap:6px;justify-content:center;border-bottom:1px solid var(--line);padding:0 10px}}.gbars i{{display:block;width:34px;border-radius:6px 6px 0 0;min-height:2px}}.gbars i.t{{background:#ffb4a8}}.gbars i.b{{background:#8fe3b0}}.gl{{display:block;text-align:center;margin-top:8px;font-size:12px;color:var(--muted)}}.gl b{{display:block;font-size:20px}}
.scatter{{margin:18px 0 26px;padding:18px 20px;background:#fff;border:1px solid var(--line);border-radius:12px}}.scatter svg{{width:100%;height:auto;max-width:640px;display:block}}
.bars.opps .bar-row{{grid-template-columns:260px 1fr 70px 110px}}.bar-p b{{font-size:12px;letter-spacing:.04em;text-transform:uppercase}}
.matrix{{margin:18px 0 26px;padding:18px 20px;background:#fff;border:1px solid var(--line);border-radius:12px}}.mhead,.mrow{{display:grid;grid-template-columns:280px 1fr 70px;gap:14px;align-items:center;padding:6px 0;border-bottom:1px dotted var(--line);font-size:14px}}.mhead{{border-bottom:1px solid var(--line);color:var(--muted);font-size:11px}}.mcells{{display:grid;grid-template-columns:repeat(10,1fr);gap:4px}}.mcells i{{display:block;height:18px;border-radius:4px;background:var(--alt)}}.mcells i.me{{background:var(--accent);box-shadow:0 0 0 2px #fff,0 0 0 3px var(--accent)}}.mcells em{{font-style:normal;text-align:center;display:block}}.mv{{text-align:right;font-weight:700}}.mv:has(+ *){{}}
.warn,b.warn{{color:#8a5a00;font-weight:700}}.h3{{font-size:19px;margin:30px 0 8px}}.h4{{font-size:16px;margin:0}}.wrap.wide{{max-width:1200px}}.two{{display:grid;grid-template-columns:1fr 1fr;gap:32px}}.asof{{font-weight:400;color:var(--muted);font-size:13px;margin-left:8px}}
@media print{{.top,.afilter,.open-build{{display:none!important}}section{{padding:28px 0;break-inside:avoid}}.hero{{background:var(--navy)!important;-webkit-print-color-adjust:exact;print-color-adjust:exact}}.pg{{break-inside:avoid}}.pglist{{display:block}}body{{font-size:13px}}h1{{font-size:32px}}h2{{font-size:22px}}}}
@media(max-width:900px){{.stiles,.stiles.six,.steps,.qgrid,.finding dl{{grid-template-columns:1fr 1fr}}.hero-grid{{grid-template-columns:1fr}}.bar-row,.bars.opps .bar-row{{grid-template-columns:1fr 60px 90px}}.bar-t{{grid-column:1/-1}}.donuts{{grid-template-columns:1fr}}.grow{{gap:8px}}.gbars i{{width:22px}}.mhead,.mrow{{grid-template-columns:1fr 60px}}.mcells{{grid-column:1/-1}}}}@media(max-width:600px){{.stiles,.stiles.six,.steps,.qgrid,.finding dl{{grid-template-columns:1fr}}.finding,.levi{{grid-template-columns:40px 1fr}}.cmpbox{{padding:20px 16px}}.cmp-row{{grid-template-columns:1fr 1fr;gap:4px}}.cmp-row.head{{display:none}}.cmp-row .lbl{{grid-column:1/-1}}.cmp-row .arrow{{display:none}}.cmp-row .today{{text-align:left}}.cmp-row .today::before{{content:"Today";display:block;font-size:11px;letter-spacing:.06em;text-transform:uppercase;color:rgba(255,255,255,.65)}}.cmp-row .build::before{{content:"The build";display:block;font-size:11px;letter-spacing:.06em;text-transform:uppercase;color:rgba(255,255,255,.65)}}}}.cmp-row>span{{min-width:0;overflow-wrap:anywhere}}
{PKG_CSS}</style></head>
<body data-first-dir="{E(dslugs[0]) if dslugs else ""}">
<header class="top">{'<span>Search and AI answer analysis</span>' if solo else '<a href="index.html">&larr; Back to the three directions</a>'}<span>Quantum Business Solutions for {E(client)}</span></header><main>
<section class="hero"><div class="wrap"><div class="hero-grid"><div><p class="eyebrow">Search and AI answers, measured {E(date)}</p><h1>{E(h["heading"])}</h1><p class="lead">{E(h["intro"])}</p></div><div class="gauge-wrap">{gauge}<p>{"Readiness on sixteen checks, averaged over every page in the sitemap." if solo else "Readiness on sixteen checks, every page: the site today against the build, scored by the same script."}</p></div></div>{_stiles(h["tiles"])}</div></section>
{('<section class="donut-band"><div class="wrap">' + donuts + '</div></section>') if donuts else ""}
<section><div class="wrap"><p class="eyebrow">Top findings</p><h2>Ranked by what they cost, with the fix</h2><p class="lead">Each finding carries its evidence, what it costs today, and what {"we would do" if solo else "the build or the 90-day plan does"} about it.</p><div class="chips">{chips}</div><div class="findings">{_finding_cards(fnd)}</div></div></section>
<section><div class="wrap"><p class="eyebrow">Why the competitors are winning</p><h2>The same pages, side by side</h2><p class="lead">{E(cmp["intro"])}</p>{cmp_tbl}<div class="verdict">{E(cmp["verdict"])}</div></div></section>
<section><div class="wrap"><p class="eyebrow">Highest leverage moves</p><h2>Five moves, in the order we would make them</h2><p class="lead">The five moves with the biggest return for the least effort, in the order we would do them. Each one is either in the build or in the first 30 days of the plan.</p><div class="lev">{moves}</div></div></section>
<section><div class="wrap"><p class="eyebrow">Whole site, every page</p><h2>{n} pages, and what each one tells a crawler</h2><p class="lead">Every URL in the sitemap, fetched {E(date)} and parsed for title, meta description, headings, words, images, links and JSON-LD schema. Each page gets a readiness score out of 100 across sixteen checks, weighted toward what answer engines read: FAQ schema, Service schema, question-form headings and depth.</p>{site_tiles}
{cmp_section}
{grades}{bytype}<h3>By section of the site</h3>{sec_tbl}<h3>Every service, industry and location page, worst first</h3>{worst_tbl}<p class="fine">Word counts exclude navigation, header and footer. The full sheet, all {n} pages and every check, is <a href="seo-audit-pages.xlsx">seo-audit-pages.xlsx</a> (also as <a href="seo-audit-pages.csv">CSV</a>){redir_note}.</p></div></section>
{every}
{blog_html}
{ent_html}
<section><div class="wrap"><p class="eyebrow">Where the traffic comes from</p><h2>Today, in numbers</h2><ul>{today_li}</ul></div></section>
<section><div class="wrap"><p class="eyebrow">Competitors</p><h2>Who earns the visits you should be earning</h2><p class="lead">{E(seo["competitors"]["intro"])}</p>{comp_bars}{comp_tbl}</div></section>
<section><div class="wrap"><p class="eyebrow">Keyword opportunities</p><h2>The terms worth a page each</h2><p class="lead">{E(seo["opportunities"]["intro"])}</p>{opps}{opp_tbl}</div></section>
<section><div class="wrap"><p class="eyebrow">Keyword clusters by service line</p><h2>What each page has to answer</h2><p class="lead">Each cluster becomes one page plus its FAQ. The terms are the headings. Volume, difficulty, cost per click and intent from Semrush, US database.</p>{clusters}</div></section>
<section><div class="wrap"><p class="eyebrow">Questions people ask</p><h2>The answers AI assistants and Google both want</h2><p class="lead">Each question becomes an FAQ item marked up as FAQPage, answered in the first sentence.</p><div class="qgrid">{qcols}</div></div></section>
<section><div class="wrap"><p class="eyebrow">Page audit</p><h2>What the pages say to a crawler today</h2>{crawl_tbl}<h3>The same pages at the competitors</h3>{cp_tbl}</div></section>
<section><div class="wrap"><p class="eyebrow">AI answer visibility</p><h2>Who is named when someone asks</h2><p class="lead">{E(seo["aeo"]["intro"])}</p>{matrix}{aeo_tbl}<h3>AI crawlers in robots.txt</h3><ul>{crawlers}</ul></div></section>
<section><div class="wrap"><p class="eyebrow">Backlinks</p><h2>{E(bl["heading"])}</h2>{auth_bars}{bl_tbl}<h3>Strongest referring domains</h3>{ref_tbl}</div></section>
<section><div class="wrap"><p class="eyebrow">Local and reviews</p><h2>What a buyer sees next to the map</h2><ul>{local_li}</ul></div></section>
<section><div class="wrap"><p class="eyebrow">The 90-day plan</p><h2>What happens, in order</h2><div class="steps">{plan}</div></div></section>
<section><div class="wrap"><p class="eyebrow">What we measure</p><h2>The same report, re-run at 30, 60 and 90 days</h2>{measure_tbl}<h3>In one list</h3><ol>{one_list}</ol><p class="note">Every number here came from a Semrush or Firecrawl pull, or from the live HTML of {E(dom)}, on {E(date)}. Where a pull was not made, the table says so. Rankings and traffic are not promised; the inputs are, and they are re-measured against this baseline.</p></div></section>
</main></body></html>'''


def _every_section(sec_id, audit, seo, build_by_path, mapping, cmp_rows, roles_avg, dom, date):
    """The every-page section: tiles, today against the build, site-wide checks, filters and one card per URL. Shared by the hub and the report."""
    s = audit["summary"]
    n = audit["count"]
    eight = _stiles([(s["avg_score"], "average readiness score out of 100"), (f"{s['grade_b_plus']} of {n}", "pages scoring B or better"), (f"{s['grade_d_f']} of {n}", "pages scoring D or F"),
                     (f"{s['faq']} of {n}", "pages with FAQPage schema"), (f"{s['service_or_article']} of {n}", "pages with Service or Article schema"), (s["review"], "pages with Review or rating schema"),
                     (f"{s['question']} of {n}", "pages with a question-form heading"), (f"{s['under_300']} of {n}", "pages under 300 words")], "eight")
    cmp_box = "".join(f'<div class="cmp-row"><span class="lbl">{E(l)}</span><span class="today">{E(t)}</span><span class="arrow" aria-hidden="true"></span><span class="build">{E(bd)}</span></div>' for l, t, bd in cmp_rows)
    opts = "".join(f'<span class="opt"><b>{E(d.title())}</b> {E(v)}</span>' for d, v in roles_avg)
    solo = build_by_path is None  # analysis only, no build to compare against
    kinds = {"rebuilt": 0, "merged": 0, "retired": 0, "migrated": 0}
    for _, k in (mapping or {}).values():
        kinds[k] += 1
    migrated = f", {kinds['migrated']} posts migrate to the blog at launch under their service line" if kinds["migrated"] else ""
    measured = []
    site = audit.get("site") or {}
    if site:
        measured = [(site.get("http_redirects_to_https", False), "http redirects to https (measured)"),
                    (site.get("one_canonical_host", False), f"One canonical host ({site.get('canonical_host') or 'two hosts answer'}) (measured)"),
                    (all(site.get(f"allows_{b}", False) for b in ("gptbot", "claudebot", "perplexitybot", "google-extended")), "robots.txt allows GPTBot, ClaudeBot, PerplexityBot and Google-Extended (measured)"),
                    (site.get("sitemap_declared", False), f"Sitemap declared in robots.txt, {site.get('sitemap_urls', 0)} URLs (measured)"),
                    (site.get("llms_txt", False), "llms.txt served (measured)"),
                    ("crawl_delay" not in site, "No crawl delay in robots.txt (measured)")]
    # authored items that restate a measured one are dropped, so the list never says the same thing twice
    dup = re.compile(r"https|canonical host|llms\.txt|gptbot|claudebot|sitemap|crawl delay", re.I)
    authored = [(ok, t) for ok, t in seo["sitewide"] if not (measured and dup.search(t))]
    sitechk = "".join(f'<li class="{"ok" if ok else "bad"}"><b>{"Yes" if ok else "No"}</b> {E(t)}</li>' for ok, t in measured + authored)
    types = {}
    for p in audit["pages"]:
        t = types.setdefault(p["type"], [0, 0]); t[0] += 1; t[1] += p["score"]
    type_opts = "".join(f'<option value="{E(k)}">{E(k.title())} pages ({v[0]}, average {round(v[1] / v[0])})</option>' for k, v in sorted(types.items(), key=lambda kv: -kv[1][0]))

    cards = []
    for p in sorted(audit["pages"], key=lambda x: x["score"]):
        tgt, kind = mapping[p["path"]] if mapping else ("", "rebuilt")
        bp = build_by_path.get("/" + tgt) if build_by_path else None
        schs = "".join(f'<span class="sch">{E(t)}</span>' for t in p["schema"] if t in ("FAQPage", "Service", "LocalBusiness", "Organization", "BlogPosting", "Article", "Review", "AggregateRating", "VideoObject", "Product"))
        checks = []
        fixed = 0
        for k, label, _w in A.CHECKS:
            tv = p["values"][k]; tok = p["checks"][k]
            if bp:
                bv = bp["values"][k]; bok = bp["checks"][k]
            else:
                bv, bok = "", False
            if not tok and bok:
                fixed += 1
            if solo:
                checks.append(f'<span class="pl1"><i>{E(label)}</i><b class="{"ok" if tok else ("warn" if str(tv).startswith("shown") else "bad")}">{E(tv)}</b></span>')
            else:
                checks.append(f'<span class="pl2"><i>{E(label)}</i><b class="{"ok" if tok else "bad"}">{E(tv)}</b><em aria-hidden="true"></em><b class="{"ok" if bok else "bad"}">{E(bv)}</b></span>')
        did = list(p["recommendations"]) if solo else [DONE[k] for k in p["failed"] if bp and bp["checks"].get(k)]
        if solo and p["type"] == "form":  # a booking or request form is not a page to rank; say so instead of asking for 300 words
            did = ["Take it out of the index (noindex) and out of the sitemap; keep one booking page per service and redirect the duplicate"] + [r for r in did if r.startswith(("Add the city", "Write a meta", "Give the page"))]
        if kind == "retired":
            did = []
        if kind == "merged":
            did.insert(0, f"Merged into {tgt}, which carries this page's earning terms")
        elif kind == "migrated":
            did.insert(0, "Migrated to the blog at launch under its service line, with BlogPosting schema, a named author and a FAQ; the post template below is the model")
        elif kind == "retired":
            did.insert(0, f"Retired; a permanent redirect sends the URL to {tgt}")
        did_li = "".join(f"<li>{E(x)}</li>" for x in did) or "<li>Nothing to fix on this page</li>"
        bscore = f'<b style="--g:{GRADE[bp["grade"]]}">{bp["score"]}<i>{bp["grade"]}</i></b><span>{fixed} checks fixed</span>' if bp else '<b style="--g:#5b616b">&ndash;</b><span>no page</span>'
        q = E(f"{p['path']} {p['title']} {p['h1']} {p['type']}".lower())
        live_a = f'<a href="https://{E(dom)}{E(p["path"])}" target="_blank" rel="noopener">{E(p["path"] if p["path"] != "/" else "/ (home page)")}</a>'
        if solo:
            cards.append(f'''<article class="pg solo" data-q="{q}" data-issues="{p['issues']}" data-type="{E(p['type'])}" data-score="{p['score']}" data-words="{p['words']}" data-url="{E(p['path'])}">
<div class="pg-id">{live_a}<small>{E(p['title'])}</small><span class="type">{E(p['type'])} page</span><div class="schs">{schs}</div></div>
<div class="pg-scores"><div class="pg-score"><span class="k">Today</span><b style="--g:{GRADE[p['grade']]}">{p['score']}<i>{p['grade']}</i></b><span>of 100, {p['issues']} of 16 checks failing</span></div></div>
<div class="pg-checks solo"><div class="legend"><span>Sixteen checks, <b class="ok">passing</b> and <b class="bad">failing</b> today</span></div>{"".join(checks)}</div>
<div class="pg-new"><h4>What we would do on this page</h4><ul class="did">{did_li}</ul><a href="https://{E(dom)}{E(p['path'])}" class="open-build" target="_blank" rel="noopener">Open the live page</a></div></article>''')
            continue
        cards.append(f'''<article class="pg" data-q="{q}" data-issues="{p['issues']}" data-type="{E(p['type'])}" data-score="{p['score']}" data-words="{p['words']}" data-url="{E(p['path'])}">
<div class="pg-id">{live_a}<small>{E(p['title'])}</small><span class="type">{E(p['type'])} page</span><div class="schs">{schs}</div></div>
<div class="pg-scores"><div class="pg-score"><span class="k">Today</span><b style="--g:{GRADE[p['grade']]}">{p['score']}<i>{p['grade']}</i></b><span>of 100</span></div><div class="pg-score build"><span class="k">In the build</span>{bscore}</div></div>
<div class="pg-checks"><div class="legend"><span>Each check: <b class="bad">today</b> <em aria-hidden="true"></em> <b class="ok">the build</b></span></div>{"".join(checks)}</div>
<div class="pg-new"><h4>What we did on this page</h4><ul class="did">{did_li}</ul><a href="{E(tgt)}" data-rel="{E(tgt)}" class="open-build" target="_blank" rel="noopener">Open the rebuilt page</a></div></article>''')

    h2 = f"{n} pages today, and what we would do on each" if solo else f"{n} pages today, and the page that replaces each one"
    cmp_block = "" if solo else f'''<div class="cmpbox"><div class="cmphead"><h3 class="h3">The same sixteen checks, run on the build</h3><p>Every recommendation on the cards below is already in place on all three directions. The scores on the right are the build, scored by the same script.</p></div><div class="cmp-rows"><div class="cmp-row head"><span class="lbl"></span><span class="today">{E(dom)} today</span><span class="arrow"></span><span class="build">The build, any direction</span></div>{cmp_box}</div><div class="opts">{opts}</div>
<p class="fine" style="margin-top:14px">Page for page: {kinds["rebuilt"]} URLs on {E(dom)} map to a rebuilt page, {kinds["merged"]} merge into a stronger page that carries their terms{migrated}, and {kinds["retired"]} retired pages redirect to the nearest page. The 301 map is in the download.</p>
<p class="fine">What is deliberately not marked up: Review and AggregateRating. Those wait for the Google review program, because marking up self-published testimonials is against Google&#39;s guidelines and we will not do it. That is the gap between the build and 100.</p></div>'''
    downloads = '<a class="btn small" href="seo-audit-pages.xlsx" download>Download the spreadsheet (Excel)</a><a class="btn small ghost" href="seo-audit-pages.csv" download>CSV</a>' + ("" if solo else '<a class="btn small ghost" href="redirects.csv" download>Redirect map (CSV)</a>')
    audit_html = f'''<section id="{sec_id}"><div class="wrap wide"><p class="eyebrow">Every page on {E(dom)}</p><h2>{h2}</h2><p class="lead">Every URL in the sitemap, fetched {E(date)} and parsed for title, meta description, headings, words, images, links and JSON-LD schema. Each page gets a readiness score out of 100 across sixteen checks, weighted toward what answer engines read: FAQ schema, Service schema, question-form headings and depth. Red is a check the page fails today. {"What we would do about each one is on the right, and every path opens the live page in a new tab." if solo else "The fixes for each page are on the right."}</p>{eight}
{cmp_block}
<h3 class="h3" style="margin-top:30px">Site-wide checks</h3><ul class="sitechk">{sitechk}</ul>
<div class="afilter"><input type="search" id="aq" placeholder="Filter by URL, title or H1" aria-label="Filter pages"><select id="atype" aria-label="Page type"><option value="">All page types</option>{type_opts}</select><select id="asort" aria-label="Sort"><option value="score-asc">Lowest score first</option><option value="score-desc">Highest score first</option><option value="words">Fewest words first</option><option value="url">By URL</option></select><label><input type="checkbox" id="aprob"> Only pages failing five or more checks</label><span id="acount" aria-live="polite">{n} pages</span>{downloads}</div>
<div class="pglist" id="alist">{"".join(cards)}</div></div>
<script>(function(){{var w=document.getElementById("alist");if(!w)return;var cards=[].slice.call(w.querySelectorAll(".pg")),q=document.getElementById("aq"),pb=document.getElementById("aprob"),ty=document.getElementById("atype"),so=document.getElementById("asort"),c=document.getElementById("acount");
var dir=(document.querySelector(".alt[href$='/index.html']")||{{}}).getAttribute?null:null;var first=document.querySelector(".pv-dirs a, .alt")?null:null;
function f(){{var s=q.value.trim().toLowerCase(),p=pb.checked,tv=ty.value,n=0;cards.forEach(function(r){{var ok=(!s||r.getAttribute("data-q").indexOf(s)>-1)&&(!p||+r.getAttribute("data-issues")>=5)&&(!tv||r.getAttribute("data-type")===tv);r.hidden=!ok;if(ok)n++}});c.textContent=n+" of "+cards.length+" pages"}}
function srt(){{var v=so.value;cards.sort(function(a,b){{if(v==="score-asc")return +a.getAttribute("data-score")-+b.getAttribute("data-score");if(v==="score-desc")return +b.getAttribute("data-score")-+a.getAttribute("data-score");if(v==="words")return +a.getAttribute("data-words")-+b.getAttribute("data-words");return a.getAttribute("data-url").localeCompare(b.getAttribute("data-url"))}});cards.forEach(function(r){{w.appendChild(r)}})}}
q.addEventListener("input",f);pb.addEventListener("change",f);ty.addEventListener("change",f);so.addEventListener("change",srt);srt();
var d0=document.body.getAttribute("data-first-dir")||"";w.querySelectorAll("a.open-build").forEach(function(a){{a.setAttribute("href",d0+"/"+a.getAttribute("data-rel"))}})}})();</script></section>'''

    return audit_html


PKG_CSS = '''
.stiles.eight{grid-template-columns:repeat(4,1fr)}
.fnds{display:grid;gap:12px;margin-top:16px}.fnd{display:grid;grid-template-columns:48px 1fr;gap:14px;background:var(--bg);border:1px solid var(--border);border-left:5px solid var(--accent);border-radius:12px;padding:18px 20px}.fnd .rank{font-size:30px;font-weight:800;color:var(--ink);line-height:1}.fnd h3{margin:6px 0 8px;font-size:19px}.fnd dl{margin:0;display:grid;gap:6px}.fnd dt{font-size:11.5px;letter-spacing:.1em;text-transform:uppercase;color:var(--muted);font-weight:700}.fnd dd{margin:2px 0 0;font-size:14.5px}
.sev{display:inline-block;font-size:12px;font-weight:700;letter-spacing:.06em;text-transform:uppercase;color:#fff;background:var(--sev);border-radius:999px;padding:3px 9px}
.wrap.wide{max-width:1200px}
.cmpbox{margin-top:28px;background:var(--chrome);color:#fff;border-radius:16px;padding:26px 28px}.cmpbox .h3,.cmpbox h3{color:#fff;margin:0 0 6px}.cmpbox p{color:rgba(255,255,255,.8);margin:0 0 16px}.cmpbox .fine{color:rgba(255,255,255,.75)}
.cmp-rows{display:grid;gap:6px}.cmp-row{display:grid;grid-template-columns:1.6fr 1fr 40px 1fr;align-items:center;gap:10px;padding:9px 12px;border-radius:10px;background:rgba(255,255,255,.06)}.cmp-row.head{background:none;font-size:12px;letter-spacing:.08em;text-transform:uppercase;color:rgba(255,255,255,.7);padding-bottom:0}.cmp-row .lbl{font-weight:600}.cmp-row .today{color:#ffb4a8;font-weight:700;text-align:right}.cmp-row .build{color:#8fe3b0;font-weight:800;font-size:17px}.cmp-row .arrow{width:40px;height:2px;background:rgba(255,255,255,.35);justify-self:center;position:relative}.cmp-row .arrow::after{content:"";position:absolute;right:-1px;top:-4px;border:5px solid transparent;border-left-color:rgba(255,255,255,.35)}.cmp-row.head .arrow{background:none}.cmp-row.head .arrow::after{display:none}
.opts{display:flex;gap:10px;flex-wrap:wrap;margin-top:14px}.opt{background:rgba(255,255,255,.1);border-radius:999px;padding:6px 12px;font-size:13.5px}.opt b{margin-right:4px}
.sitechk{list-style:none;padding:0;margin:8px 0 0;display:grid;grid-template-columns:1fr 1fr;gap:6px 18px}.sitechk li{padding:8px 12px;border:1px solid var(--border);border-radius:10px;font-size:14.5px;background:var(--bg)}.sitechk li b{display:inline-block;min-width:34px;font-weight:800}.sitechk li.ok b{color:#1b7f4b}.sitechk li.bad b{color:#b3261e}
.afilter{display:flex;gap:10px;flex-wrap:wrap;align-items:center;margin:26px 0 14px}.afilter input[type=search],.afilter select{font:inherit;font-size:14.5px;padding:9px 12px;border:1px solid var(--border);border-radius:10px;background:var(--bg);color:var(--fg);min-height:44px}.afilter input[type=search]{flex:1;min-width:220px}.afilter label{font-size:14px;display:inline-flex;align-items:center;gap:6px}.afilter #acount{color:var(--muted);font-size:14px;margin-left:auto}.btn.small{padding:9px 14px;font-size:14px}.btn.ghost{background:transparent;border:1px solid var(--accent);color:var(--fg)}
.pglist{display:grid;gap:12px}.pg{display:grid;grid-template-columns:1.1fr .8fr 2fr 1.2fr;gap:18px;padding:16px 18px;border:1px solid var(--border);border-radius:14px;background:var(--bg)}.pg[hidden]{display:none}
.pg-id a{color:var(--ink);font-weight:700;text-decoration:none;word-break:break-all}.pg-id small{display:block;color:var(--muted);font-size:12.5px;margin-top:2px}.pg-id .type{display:inline-block;margin-top:8px;font-size:11.5px;letter-spacing:.08em;text-transform:uppercase;color:var(--muted);font-weight:700}.schs{display:flex;gap:4px;flex-wrap:wrap;margin-top:8px}.sch{font-size:11.5px;border:1px solid var(--border);border-radius:999px;padding:2px 8px;color:var(--muted)}
.pg-scores{display:grid;gap:10px;align-content:start}.pg-score .k{display:block;font-size:11.5px;letter-spacing:.08em;text-transform:uppercase;color:var(--muted);font-weight:700}.pg-score b{display:inline-block;font-size:30px;line-height:1;color:var(--g);letter-spacing:-.02em}.pg-score b i{font-style:normal;font-size:14px;margin-left:4px;color:#fff;background:var(--g);border-radius:6px;padding:1px 6px;vertical-align:middle}.pg-score span:last-child{display:block;font-size:12.5px;color:var(--muted)}.pg-score.build{padding-top:6px;border-top:1px dashed var(--border)}
.pg-checks{display:grid;grid-template-columns:1fr 1fr;gap:4px 14px;font-size:12.5px;align-content:start}.pg-checks .legend{grid-column:1/-1;color:var(--muted);font-size:12px;margin-bottom:2px}.pl2{display:grid;grid-template-columns:1fr auto 14px auto;gap:6px;align-items:center;padding:3px 0;border-bottom:1px dotted var(--border)}.pl2 i{font-style:normal;color:var(--muted)}.pl2 b.ok{color:#1b7f4b}.pl2 b.bad{color:#b3261e}.pl2 em,.legend em{display:inline-block;width:12px;height:2px;background:var(--muted);position:relative}.pl2 em::after,.legend em::after{content:"";position:absolute;right:-2px;top:-3px;border:4px solid transparent;border-left-color:var(--muted)}
.pg-new h4{margin:0 0 6px;font-size:12px;letter-spacing:.08em;text-transform:uppercase;color:var(--muted)}.pg-new .did{margin:0;padding-left:16px;font-size:13.5px}.pg-new .did li{margin:2px 0}.pg-new .open-build{display:inline-block;margin-top:8px;font-size:13.5px;font-weight:600;color:var(--ink)}
.cmp-row>span{min-width:0;overflow-wrap:anywhere}
.pg.solo{grid-template-columns:1.2fr .7fr 2fr 1.3fr}.pl1{display:grid;grid-template-columns:1fr auto;gap:8px;padding:2px 0;border-bottom:1px dotted var(--border)}.pl1 i{font-style:normal;color:var(--muted)}.pl1 b{font-weight:700;text-align:right}
b.warn,.warn{color:#8a5a00;font-weight:700}
@media(max-width:1000px){.pg,.pg.solo{grid-template-columns:1fr 1fr}.pg-checks{grid-column:1/-1}.stiles.eight{grid-template-columns:repeat(2,1fr)}}@media(max-width:600px){.pg,.pg.solo,.pg-checks,.sitechk{grid-template-columns:1fr}.cmpbox{padding:20px 16px}.cmp-row{grid-template-columns:1fr 1fr;gap:4px}.cmp-row.head{display:none}.cmp-row .lbl{grid-column:1/-1}.cmp-row .arrow{display:none}.cmp-row .today{text-align:left}.cmp-row .today::before{content:"Today ";font-size:11px;letter-spacing:.06em;text-transform:uppercase;color:rgba(255,255,255,.65);display:block}.cmp-row .build::before{content:"The build ";font-size:11px;letter-spacing:.06em;text-transform:uppercase;color:rgba(255,255,255,.65);display:block;font-weight:600}}
'''


def hub_fragments(content, seo, audit, build_by_path, mapping, cmp_rows, roles_avg, dom, date):
    client = content["client"]
    short = client.split(" ")[0]
    h = seo["hub"]
    tiles = _stiles(seo["hero"]["tiles"])
    top = _finding_cards(seo["findings"][:5], "fnd")
    today = "".join(f"<li>{E(x)}</li>" for x in h["today"])
    after = "".join(f"<li>{E(x)}</li>" for x in h["after"])
    keep = "".join(f"<tr><td><code>{E(o)}</code></td><td><code>{E(n)}</code></td><td>{E(w)}</td></tr>" for o, n, w in h.get("keep", []))
    keep_tbl = (f'<h3 class="h3" style="margin-top:34px">{E(h.get("keep_heading", "Every page that earns a visitor today keeps earning it"))}</h3><div class="tblwrap" tabindex="0"><table class="seo"><thead><tr><th scope="col">Today</th><th scope="col">After launch (301)</th><th scope="col">What it ranks for now</th></tr></thead><tbody>{keep}</tbody></table></div>' if keep else "")
    search = f'''<section id="search"><div class="wrap"><p class="eyebrow">Search and AI answers</p><h2>{E(h["heading"])}</h2><p class="lead">{E(h["intro"])}</p>
{tiles}
<h3 class="h3" style="margin-top:34px">Top findings, ranked by what they cost</h3><div class="fnds">{top}</div>
<div class="two" style="margin-top:34px"><div><h3 class="h3">Today <span class="asof">Semrush US database and live HTML, {E(date)}</span></h3><ul>{today}</ul></div><div><h3 class="h3">After launch</h3><ul>{after}</ul></div></div>
{keep_tbl}
<p style="margin-top:26px"><a class="btn" href="seo-report.html" target="_blank" rel="noopener">Read the full SEO and AI search analysis</a></p>
<div class="change" style="margin-top:22px"><strong>What we do not promise:</strong> {E(h.get("note", "Rankings or traffic. Those depend on the market and on what you publish after launch. What we promise is that nothing about the build is the reason they do not come, that every number above is re-measured at 30, 60 and 90 days, and that you see the same report we do."))}</div></div></section>'''

    audit_html = _every_section("audit", audit, seo, build_by_path, mapping, cmp_rows, roles_avg, dom, date)
    css = PKG_CSS
    return {"search": search, "audit": audit_html, "css": css, "nav": '<a href="#audit">Every URL</a>'}


def to_xlsx(audit, seo, mapping, path):
    """Every table in the analysis as one workbook: every page with every check, findings, opportunities, questions,
    answer-engine results, competitors, the redirect map when there is a build, and the measured site checks."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF"); head_fill = PatternFill("solid", fgColor="0F1E33")
    grade_fill = {"A": "DDF3E4", "B": "EAF5DC", "C": "FFF1CC", "D": "FFE2D6", "F": "FFD2D2"}

    def sheet(title, head, rows, widths=None, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title[:31]
        ws.append(head)
        for c in ws[1]:
            c.font = head_font; c.fill = head_fill; c.alignment = Alignment(vertical="center", wrap_text=True)
        for r in rows:
            ws.append([("" if v is None else v) for v in r])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, w in enumerate(widths or [], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        return ws

    yn = lambda v: "yes" if v else "no"
    head = ["URL", "Path", "Page type", "Score", "Grade", "Checks failing", "Title", "Title length", "City in title", "Meta description", "Meta length", "H1", "H1 count", "H2 count",
            "Question headings", "Words", "Images", "Images missing alt", "Internal links", "Canonical", "FAQPage schema", "FAQ shown on page", "Service or Article schema", "LocalBusiness schema", "Review schema", "Schema types", "Noindex", "Form on page", "What we would do"]
    if mapping:
        head += ["New path in the build", "Kind"]
    rows = []
    for p in sorted(audit["pages"], key=lambda x: x["score"]):
        r = [p["url"], p["path"], p["type"], p["score"], p["grade"], p["issues"], p["title"], p["title_len"], yn(p["checks"]["city"]), p["meta"], p["meta_len"], p["h1"], p["h1_count"], p["h2_count"],
             p["question_headings"], p["words"], p["images"], p["images_no_alt"], p["internal_links"], p["canonical"], yn(p["faq"]), yn(p.get("faq_visible")), yn(p["service"] or p["article"]), yn(p["local"]), yn(p["review"]),
             ", ".join(p["schema"]), yn(p["noindex"]), yn(p["has_form"]), "; ".join(p["recommendations"])]
        if mapping:
            tgt, kind = mapping[p["path"]]
            r += ["/" + tgt, kind]
        rows.append(r)
    ws = sheet("Every page", head, rows, [48, 30, 10, 7, 7, 9, 48, 8, 8, 60, 8, 40, 6, 6, 8, 8, 7, 8, 8, 40, 9, 9, 10, 10, 9, 40, 8, 8, 90] + ([36, 10] if mapping else []), first=True)
    for row in ws.iter_rows(min_row=2):
        g = row[4].value
        if g in grade_fill:
            row[3].fill = row[4].fill = PatternFill("solid", fgColor=grade_fill[g])
    sheet("Findings", ["Rank", "Severity", "Finding", "Evidence", "What it costs", "What we do"], [[i, f["severity"], f["title"], f.get("evidence", ""), f.get("cost", ""), f.get("fix", "")] for i, f in enumerate(seo["findings"], 1)], [6, 10, 48, 90, 60, 70])
    sheet("Moves", ["Order", "Move", "When", "Where", "Detail"], [[i, m["title"], m["when"], m["where"], m["body"]] for i, m in enumerate(seo["moves"], 1)], [6, 48, 18, 16, 110])
    sheet("Opportunities", ["Keyword", "Searches a month", "Difficulty", "Position today", "Who ranks"], [list(r) for r in seo["opportunities"]["rows"]], [40, 14, 10, 16, 70])
    sheet("Questions", ["Topic", "Question", "Searches a month"], [[q["topic"], it[0], it[1]] for q in seo["questions"] for it in q["items"]], [28, 70, 14])
    sheet("AI answers", ["Query", "Position"] + [f"#{i}" for i in range(1, 11)], [[r["query"], r.get("position") or "absent"] + list(r["results"][:10]) + [""] * (10 - len(r["results"][:10])) for r in seo["aeo"]["rows"]], [40, 9] + [26] * 10)
    sheet("Competitors", ["Domain", "Organic visits a month", "Organic keywords", "Authority Score"], [list(r) for r in seo["competitors"]["rows"]], [40, 16, 14, 12])
    sheet("Backlinks", ["Domain", "Authority Score", "Referring domains", "Backlinks"], [list(r) for r in seo["backlinks"]["rows"]], [40, 12, 14, 20])
    site = audit.get("site") or {}
    sheet("Site checks", ["Check", "Result"], [[k, ("yes" if v is True else "no" if v is False else v)] for k, v in site.items()] + [[t, "yes" if ok else "no"] for ok, t in seo["sitewide"]], [70, 30])
    if mapping:
        sheet("Redirects", ["Old URL", "New path in the build", "Status", "Kind"], [[p["path"], "/" + mapping[p["path"]][0], 301, mapping[p["path"]][1]] for p in audit["pages"]], [50, 44, 8, 10])
    wb.save(path)


def write_analysis(slug, cfg, out_dir):
    """Analysis only, for a prospect with no build yet: the same report, every-page cards, CSV and workbook, without the build column.
    cfg: {"client", "domain", "accent", "chrome", "cities"}; seo and audit come from brands/<slug>.seo.json and .audit.json."""
    root = os.path.dirname(HERE)
    content_path = os.path.join(root, "brands", f"{slug}.content.json")
    seo, audit = _load(content_path)
    if not seo:
        sys.exit(f"need brands/{slug}.seo.json and brands/{slug}.audit.json")
    validate(seo, content_path.replace(".content.json", ".seo.json"))
    content = {"client": cfg["client"], "brand": {"accent": cfg.get("accent", "#2f6f4e"), "chrome_bg": cfg.get("chrome", "#0f1e33")}, "schema": {"cities": cfg.get("cities", [])}}
    os.makedirs(out_dir, exist_ok=True)
    open(os.path.join(out_dir, "seo-report.html"), "w", encoding="utf-8").write(report_html(content, seo, audit, None, [], cfg.get("base_url", ""), [], [], None, None))
    A.to_csv(audit, os.path.join(out_dir, "seo-audit-pages.csv"))
    to_xlsx(audit, seo, None, os.path.join(out_dir, "seo-audit-pages.xlsx"))
    print(f"  analysis: report, {audit['count']}-page audit, CSV and workbook in {out_dir}")


def write(content, content_path, themes, roles, base, out_dir, slug_of):
    seo, audit = _load(content_path)
    if not seo:
        return None
    validate(seo, content_path.replace(".content.json", ".seo.json"))
    sch = content.get("schema") or {}
    cities = ",".join(sch.get("cities", []))
    dslugs = [slug_of(t) for t in themes]
    type_rules = seo.get("build_types", [["^/services/", "service"], ["^/industries/", "industry"], ["^/products/", "service"], ["^/locations/", "city"], ["^/blog/", "post"]])
    builds = {}
    for d in dslugs:
        pages = []
        root = os.path.join(out_dir, d)
        for r_, _, files in os.walk(root):
            for fn in files:
                if fn.endswith(".html"):
                    rel = os.path.relpath(os.path.join(r_, fn), root).replace(os.sep, "/")
                    res = A.audit_html(open(os.path.join(r_, fn), "rb").read(), rel, "", [c for c in cities.split(",") if c], is_build=True)
                    if res:
                        res["path"] = "/" + rel
                        res["type"] = A.classify("/" + rel, type_rules)
                        pages.append(res)
        builds[d] = {"pages": pages, "summary": A.summarize(pages)}
    first = dslugs[0]
    build = builds[first]
    build_by_path = {p["path"]: p for p in build["pages"]}
    build_files = {p["path"].lstrip("/") for p in build["pages"]}
    mapping = map_pages(audit["pages"], build_files, seo)
    cmp_rows = _cmp_rows(audit["summary"], build["summary"], audit["count"], build["summary"]["count"])
    roles_avg = [(d, f"{builds[d]['summary']['avg_score']} average, {builds[d]['summary']['grade_a']} of {builds[d]['summary']['count']} A") for d in dslugs]
    dom = audit["domain"]
    date = seo.get("measured") or audit.get("measured", "")

    open(os.path.join(out_dir, "seo-report.html"), "w", encoding="utf-8").write(report_html(content, seo, audit, build, cmp_rows, base, dslugs, roles_avg, build_by_path, mapping))
    A.to_csv(audit, os.path.join(out_dir, "seo-audit-pages.csv"))
    to_xlsx(audit, seo, mapping, os.path.join(out_dir, "seo-audit-pages.xlsx"))
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([f"old_url_on_{dom}", "new_path_in_build", "status", "kind"])
    for p in audit["pages"]:
        tgt, kind = mapping[p["path"]]
        w.writerow([p["path"], "/" + tgt, "301", kind])
    open(os.path.join(out_dir, "redirects.csv"), "w", encoding="utf-8", newline="").write(buf.getvalue())
    json.dump({"first": first, "summaries": {d: builds[d]["summary"] for d in dslugs}}, open(os.path.join(out_dir, "seo-build-scores.json"), "w"), indent=1)
    frags = hub_fragments(content, seo, audit, build_by_path, mapping, cmp_rows, roles_avg, dom, date)
    frags["first_dir"] = first
    print(f"  seo: report, {audit['count']}-page audit, redirects; build {first} averages {build['summary']['avg_score']} ({build['summary']['grade_a']} of {build['summary']['count']} A)")
    return frags
